"""Export graded results: an Excel workbook for the professor and a CSV for the gradebook."""
from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .store import GradeStore

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")   # needs review
APPROVED_FILL = PatternFill("solid", fgColor="E2EFDA")
OVERRIDE_FONT = Font(bold=True, color="C00000")
THIN = Side(style="thin", color="BFBFBF")
WRAP = Alignment(wrap_text=True, vertical="top")


@dataclass
class ExportSummary:
    students: int
    approved: int
    needs_review: int
    not_graded: int
    path: str

    def __str__(self) -> str:
        return (f"{self.path}: {self.students} students — {self.approved} approved, "
                f"{self.needs_review} need review, {self.not_graded} not graded")


def _header(ws, values: list[str]) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[ws.max_row].height = 32


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_excel(store: GradeStore, assignment: str, path: str | Path) -> ExportSummary:
    aid, rubric = store.get_assignment(assignment)
    subs = store.submissions(aid)
    wb = Workbook()

    # ---------------- Summary ----------------
    ws = wb.active
    ws.title = "Summary"
    fixed = ["Student ID", "Name", "Email", "Status"]
    crit_headers = [f"{c.name} ({c.points:g})" for c in rubric.criteria]
    _header(ws, fixed + crit_headers + ["Total", "Out of", "Percent", "Late", "Flags"])
    first_crit_col = len(fixed) + 1
    last_crit_col = first_crit_col + len(rubric.criteria) - 1
    total_col = last_crit_col + 1
    fc, lc, tc = get_column_letter(first_crit_col), get_column_letter(last_crit_col), get_column_letter(total_col)
    oc, pc = get_column_letter(total_col + 1), get_column_letter(total_col + 2)

    counts = {"approved": 0, "graded": 0, "pending": 0, "error": 0}
    for s in subs:
        counts[s.status] = counts.get(s.status, 0) + 1
        sc = store.scores(s.id)
        r = ws.max_row + 1
        label = {"approved": "Approved", "graded": "Needs review", "pending": "Not graded", "error": "Error"}[s.status]
        row = [s.student_id, s.student_name, s.email, label]
        for c in rubric.criteria:
            row.append(sc[c.id]["final_score"] if c.id in sc else None)
        row += [
            f'=IF(COUNT({fc}{r}:{lc}{r})={len(rubric.criteria)},SUM({fc}{r}:{lc}{r}),"")',
            rubric.total_points,
            f'=IF({tc}{r}="","",{tc}{r}/{oc}{r})',
            "Yes" if s.late else "",
            "; ".join(s.flags),
        ]
        ws.append(row)
        ws.cell(r, 4).fill = APPROVED_FILL if s.status == "approved" else REVIEW_FILL
        ws.cell(r, total_col + 2).number_format = "0.0%"
        for j, c in enumerate(rubric.criteria):
            if c.id in sc and sc[c.id]["overridden"]:
                cell = ws.cell(r, first_crit_col + j)
                cell.font = OVERRIDE_FONT
    last_data = ws.max_row
    if subs:
        ws.append([])
        avg_row = ws.max_row + 1
        ws.cell(avg_row, 2, "Class average").font = Font(bold=True)
        for col in range(first_crit_col, total_col + 1):
            L = get_column_letter(col)
            cell = ws.cell(avg_row, col, f'=IFERROR(AVERAGE({L}2:{L}{last_data}),"")')
            cell.number_format = "0.0"
            cell.font = Font(bold=True)
        ws.cell(avg_row, total_col + 2, f'=IFERROR(AVERAGE({pc}2:{pc}{last_data}),"")').number_format = "0.0%"
        ws.auto_filter.ref = f"A1:{get_column_letter(total_col + 4)}{last_data}"
    ws.freeze_panes = ws.cell(2, first_crit_col)
    _widths(ws, [12, 22, 26, 14] + [13] * len(rubric.criteria) + [9, 8, 9, 6, 60])
    ws.cell(1, total_col + 4).alignment = Alignment(wrap_text=False)

    # ---------------- Detail ----------------
    wd = wb.create_sheet("Detail")
    _header(wd, ["Student ID", "Name", "Criterion", "Source", "AI level", "AI score", "Final score",
                 "Changed by professor", "AI confidence", "AI reason", "Evidence", "Professor comment"])
    for s in subs:
        sc = store.scores(s.id)
        for c in rubric.criteria:
            row = sc.get(c.id)
            if row is None:
                wd.append([s.student_id, s.student_name, c.name, c.source, None, None, None, "", None, "", "", ""])
                continue
            wd.append([
                s.student_id, s.student_name, c.name, c.source, row["ai_level"], row["ai_score"], row["final_score"],
                "Yes" if row["overridden"] else "", row["confidence"], row["ai_reason"] or "",
                ", ".join(json.loads(row["evidence"] or "[]")), row["comment"] or "",
            ])
            if row["overridden"]:
                wd.cell(wd.max_row, 7).font = OVERRIDE_FONT
            wd.cell(wd.max_row, 9).number_format = "0%"
    for r in wd.iter_rows(min_row=2):
        r[9].alignment = WRAP
        r[11].alignment = WRAP
    wd.freeze_panes = "C2"
    wd.auto_filter.ref = f"A1:L{max(wd.max_row, 1)}"
    _widths(wd, [12, 22, 24, 10, 14, 9, 10, 11, 11, 55, 22, 40])
    # Highlight low-confidence rows for review.
    wd.conditional_formatting.add(
        f"A2:L{max(wd.max_row, 2)}", FormulaRule(formula=["AND(ISNUMBER($I2),$I2<0.7)"], fill=REVIEW_FILL)
    )

    # ---------------- Feedback ----------------
    wf = wb.create_sheet("Feedback")
    _header(wf, ["Student ID", "Name", "Total", "Out of", "Feedback to student"])
    for s in subs:
        wf.append([s.student_id, s.student_name, store.total(s.id, rubric), rubric.total_points, s.feedback])
        wf.cell(wf.max_row, 5).alignment = WRAP
    _widths(wf, [12, 22, 8, 8, 90])
    wf.freeze_panes = "C2"

    # ---------------- Flags ----------------
    wl = wb.create_sheet("Flags")
    _header(wl, ["Student ID", "Name", "Issue"])
    for s in subs:
        for f in s.flags:
            wl.append([s.student_id, s.student_name, f])
    if wl.max_row == 1:
        wl.append(["", "", "No issues found"])
    _widths(wl, [12, 22, 90])

    # ---------------- Rubric ----------------
    wr = wb.create_sheet("Rubric")
    wr.append([rubric.assignment])
    wr["A1"].font = Font(bold=True, size=13)
    wr.append([f"Total: {rubric.total_points:g} points"])
    wr.append([])
    _header(wr, ["Criterion ID", "Criterion", "Graded from", "Points", "Level", "Level score", "Description"])
    for c in rubric.criteria:
        for l in c.levels:
            wr.append([c.id, c.name, c.source, c.points, l.label, l.score, l.description])
            wr.cell(wr.max_row, 7).alignment = WRAP
    _widths(wr, [18, 26, 12, 8, 14, 11, 80])

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row):
            for cell in row:
                if cell.value is not None:
                    cell.border = Border(bottom=THIN)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return ExportSummary(
        students=len(subs), approved=counts.get("approved", 0), needs_review=counts.get("graded", 0),
        not_graded=counts.get("pending", 0) + counts.get("error", 0), path=str(path),
    )


# ---------------- Gradebook CSV ----------------
ID_COLUMNS = ("sis user id", "sis login id", "id", "student_id", "student id")


def export_gradebook_csv(
    store: GradeStore,
    assignment: str,
    path: str | Path,
    lms_template: str | Path | None = None,
    column: str | None = None,
    approved_only: bool = True,
) -> str:
    """Write grades for the gradebook.

    Without a template: a simple CSV (student_id, student_name, email, score, points_possible).
    With `lms_template` (the gradebook CSV exported from the LMS): fills the
    assignment's column in that file, matching students by ID, and leaves every
    other cell untouched so the file can be imported straight back.
    Returns a one-line report.
    """
    aid, rubric = store.get_assignment(assignment)
    subs = store.submissions(aid)
    grades = {}
    skipped = 0
    for s in subs:
        if approved_only and s.status != "approved":
            skipped += 1
            continue
        total = store.total(s.id, rubric)
        if total is None:
            skipped += 1
            continue
        grades[s.student_id] = total

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if lms_template is None:
        with path.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["student_id", "student_name", "email", "score", "points_possible"])
            for s in subs:
                if s.student_id in grades:
                    w.writerow([s.student_id, s.student_name, s.email, f"{grades[s.student_id]:g}", f"{rubric.total_points:g}"])
        return f"{path}: {len(grades)} grades written, {skipped} skipped (not approved or incomplete)"

    with Path(lms_template).open(newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    if not rows:
        raise ValueError("The LMS gradebook file is empty.")
    header = rows[0]
    lower = [h.strip().lower() for h in header]
    # Assignment column: exact name given, or a header that starts with the assignment name
    # (LMS exports often append an id, e.g. "Project 2 (12345)").
    if column:
        if column not in header:
            raise ValueError(f"Column '{column}' not found in the LMS file.")
        col = header.index(column)
    else:
        target = re.sub(r"\s+", " ", assignment.strip().lower())
        matches = [i for i, h in enumerate(lower) if re.sub(r"\s+", " ", h).startswith(target)]
        if len(matches) != 1:
            raise ValueError(
                f"Found {len(matches)} columns matching '{assignment}' in the LMS file. "
                "Pass the exact column header with --column."
            )
        col = matches[0]
    id_cols = [lower.index(c) for c in ID_COLUMNS if c in lower]
    if not id_cols:
        raise ValueError("The LMS file has no student ID column (looked for: " + ", ".join(ID_COLUMNS) + ").")

    filled, matched = 0, set()
    for row in rows[1:]:
        for ic in id_cols:
            if ic < len(row) and row[ic].strip() in grades:
                sid = row[ic].strip()
                while len(row) <= col:
                    row.append("")
                row[col] = f"{grades[sid]:g}"
                matched.add(sid)
                filled += 1
                break
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)
    unmatched = sorted(set(grades) - matched)
    msg = f"{path}: filled '{header[col]}' for {filled} students, {skipped} skipped (not approved or incomplete)"
    if unmatched:
        msg += f"; not found in LMS file: {', '.join(unmatched)}"
    return msg


def export_feedback_files(store: GradeStore, assignment: str, folder: str | Path) -> int:
    """One plain-text feedback file per student (score breakdown + comments)."""
    aid, rubric = store.get_assignment(assignment)
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    n = 0
    for s in store.submissions(aid):
        sc = store.scores(s.id)
        if not sc:
            continue
        total = store.total(s.id, rubric)
        lines = [rubric.assignment, f"{s.student_name} ({s.student_id})", ""]
        for c in rubric.criteria:
            r = sc.get(c.id)
            score = "—" if r is None or r["final_score"] is None else f"{r['final_score']:g}"
            lines.append(f"{c.name}: {score} / {c.points:g}")
            if r is not None:
                note = r["comment"] or r["ai_reason"]
                if note:
                    lines.append(f"    {note}")
        lines += ["", f"Total: {'—' if total is None else f'{total:g}'} / {rubric.total_points:g}", "", s.feedback]
        safe = re.sub(r"[^A-Za-z0-9_-]+", "_", f"{s.student_id}_{s.student_name}").strip("_")
        (folder / f"{safe}.txt").write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
        n += 1
    return n
