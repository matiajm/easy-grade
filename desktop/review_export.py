"""Exports for the review screen: Valery's Excel layout (Grades, Details, Rubric) and the professor's 4-column CSV.

Both go through the export guard (easygrade/pipeline/export/guard.py): only approved ("reviewed") students are
written, and only if their total is the sum of their scores and every criterion has a score. A student the guard
refuses is listed in the result with the reason instead of being written.

Rows come from ReviewMixin.review_rows(): one dict per student/team, see desktop/review.py.
The workbook is built with openpyxl so nothing is loaded from the internet.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from easygrade.pipeline.export.csv_export import export_csv
from easygrade.pipeline.export.guard import refusal_reasons

INK, MUTED, VIOLET = "1B1830", "6A6480", "3B2F63"
STRIPE, BAND, EDIT, LINE = "F7F5FB", "EFE8FC", "FFF7D6", "E4E1EB"
OK, OK_SOFT, WARN, WARN_SOFT, GOOD_SOFT = "1D7A4A", "E3F3EA", "8F5400", "FCF1DC", "EEEDF6"


# ---------------------------------------------------------------- guard adapter
def guard_rubric(rubric) -> dict:
    """The rubric in the shape the guard reads. A level covers its floor up to just below the level above it."""
    sections = []
    for c in rubric.criteria:
        levels, ceiling = [], c.points
        for lvl in c.levels:  # best to worst
            levels.append({"level": lvl.label, "min_points": lvl.floor, "max_points": ceiling})
            ceiling = lvl.floor - 1e-9
        sections.append({"section_id": c.id, "max_points": c.points, "levels": levels})
    return {"sections": sections}


def guard_review(row: dict, rubric) -> dict:
    """A review.json-style dict for one student row, so the same rules the plan defines apply here."""
    sections = []
    for c in row["criteria"]:
        sections.append({"section_id": c["id"], "final_score": c["final"], "final_level": c["level"]})
    return {"team_id": row["id"], "students": [{"name": m} for m in row["members"]],
            "sections": sections, "total": row["total"],
            "status": "reviewed" if row["status"] == "approved" else "needs_review",
            "flags": row["check"], "ai_usage": {"suggested": None}}


def split_by_guard(rows: list[dict], rubric) -> tuple[list[dict], dict[str, list[str]]]:
    """(rows the guard allows, {student id: reasons} for the ones it refuses)."""
    g = guard_rubric(rubric)
    ok, refused = [], {}
    for row in rows:
        reasons = refusal_reasons(guard_review(row, rubric), g)
        if reasons:
            refused[row["id"]] = reasons
        else:
            ok.append(row)
    return ok, refused


# ---------------------------------------------------------------- text columns
def attention_text(row: dict) -> str:
    items = [i for i in row["check"] if "needs your score" not in i]
    entered = [c for c in row["criteria"] if c["ai_score"] is None and c["final"] is not None]
    for c in entered:
        items.append(f"You entered the {c['name']} score")
    diff = row.get("changed_by")
    if diff:
        items.append(f"You changed the suggested score by {'+' if diff > 0 else ''}{diff:g}")
    return "; ".join(items)


def ai_text(row: dict) -> str:
    ai = row.get("ai") or {}
    if ai.get("signals"):
        return "Review: " + "; ".join(ai["signals"])
    return "Not checked" if ai.get("status") == "not_checked" else "No concern"


# ---------------------------------------------------------------- csv
def export_review_csv(rows: list[dict], rubric, out_csv: Path, log_path: Path | None = None) -> dict:
    """The professor's four columns. One CSV row per team member, all with the team's grade."""
    reviews = [(r["id"], guard_review(r, rubric) | {"flags": [attention_text(r)],
                                                    "ai_usage": {"suggested": None}}) for r in rows]
    result = export_csv(reviews, guard_rubric(rubric), out_csv, log_path)
    return {"rows_written": result.rows_written, "exported": result.exported, "refused": result.refused}


# ---------------------------------------------------------------- excel
def _font(size=11, bold=False, color=INK, italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)


def _fill(argb):
    return PatternFill("solid", fgColor=argb)


def _put(ws, r, col, value, bold=False, size=11, color=INK, center=False, fill=None):
    c = ws.cell(row=r, column=col, value=value)
    c.font = _font(size, bold, color)
    c.alignment = Alignment(vertical="top", horizontal="center") if center else Alignment(vertical="top", wrap_text=True, indent=1)
    c.border = Border(bottom=Side(style="thin", color=LINE))
    if fill:
        c.fill = _fill(fill)
    return c


def _title(ws, text, subtitle):
    ws["A1"] = text
    ws["A1"].font = _font(16, True)
    ws["A2"] = subtitle
    ws["A2"].font = _font(10, False, MUTED, True)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[3].height = 8


def _header(ws, labels):
    ws.row_dimensions[4].height = 24
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=4, column=i, value=label)
        c.font = _font(11, True, "FFFFFF")
        c.fill = _fill(VIOLET)
        c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)


def _fit(ws, r, widths, texts):
    lines = max([1] + [-(-len(str(t)) // max(int(widths[i] * 1.25), 1)) for i, t in enumerate(texts) if t])
    ws.row_dimensions[r].height = max(20, 15 * lines + 6)


def _level_formula(row: int, n: int, nlevels: int, labels: list[str], first_col: int, ra: str, rbounds: list[str]) -> str:
    """Excel formula giving the level name for the score in column E, read from the Rubric tab."""
    expr = f'"{labels[-1]}"'
    for i in range(nlevels - 2, -1, -1):
        expr = f'IF(E{row}>=INDEX({rbounds[i]},MATCH(C{row},{ra},0)),"{labels[i]}",{expr})'
    return "=" + expr


def build_review_workbook(rows: list[dict], rubric, path: Path, title: str | None = None) -> Path:
    """Write Grades / Details / Rubric for the given (already guard-approved) rows."""
    wb = Workbook()
    gs, ds, rs = wb.active, wb.create_sheet("Details"), wb.create_sheet("Rubric")
    gs.title = "Grades"
    for ws in (gs, ds, rs):
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A5"
    title = title or rubric.assignment
    labels = [lv.label for lv in rubric.criteria[0].levels] if rubric.criteria else []
    nlevels = len(labels)

    # Rubric: the point ranges the formulas read. One "from" column per level except the lowest.
    _title(rs, "Rubric", f"{rubric.assignment} · {rubric.total_points:g} points. Each level starts at the score shown and runs up to the next level.")
    _header(rs, ["Category", "Points"] + [f"{l} from" for l in labels[:-1]] + [labels[-1] if labels else "Level"])
    for i, c in enumerate(rubric.criteria):
        r = 5 + i
        fill = STRIPE if r % 2 == 0 else None
        _put(rs, r, 1, c.name, fill=fill)
        _put(rs, r, 2, c.points, center=True, fill=fill)
        for j, lvl in enumerate(c.levels[:-1]):
            _put(rs, r, 3 + j, lvl.floor, center=True, fill=fill)
        last_from = c.levels[-2].floor if len(c.levels) > 1 else c.points
        _put(rs, r, 2 + nlevels, f"0 to {last_from - 1:g}" if len(c.levels) > 1 else "any", center=True, fill=fill)
    rt = 5 + len(rubric.criteria)
    _put(rs, rt, 1, "Total", bold=True)
    _put(rs, rt, 2, f"=SUM(B5:B{rt - 1})", bold=True, center=True)
    for i, w in enumerate([38, 10] + [16] * nlevels, start=1):
        rs.column_dimensions[get_column_letter(i)].width = w
    last = rt - 1
    ra = f"Rubric!$A$5:$A${last}"
    rb = f"Rubric!$B$5:$B${last}"
    rbounds = [f"Rubric!${get_column_letter(3 + j)}$5:${get_column_letter(3 + j)}${last}" for j in range(nlevels - 1)]

    # Details: a band per team, one row per category, then the team total. Column A holds the team id on
    # category rows only, so the Grades tab can add up exactly one team without matching names.
    _title(ds, "How each grade adds up", "Yellow scores can be changed. The level, the team total and the Grades tab update.")
    _header(ds, ["Team", "Students", "Category", "Level", "Score", "Out of", "Why this score", "Evidence"])
    widths = [11, 24, 35, 13, 9, 9, 55, 50]
    r = 5
    for row in rows:
        ds.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        band = ds.cell(row=r, column=1, value=f"{row['id']}  ·  {', '.join(row['members'])}")
        band.font = _font(11, True, VIOLET)
        band.fill = _fill(BAND)
        band.alignment = Alignment(vertical="center", indent=1)
        ds.row_dimensions[r].height = 22
        r += 1
        start = r
        for c in row["criteria"]:
            entered = c["ai_score"] is None
            why = c["comment"] or ("Entered by the professor." if entered else c["reason"])
            ev = "" if entered else "; ".join(c["evidence"])
            _put(ds, r, 1, row["id"], color=MUTED, size=10)
            _put(ds, r, 2, ", ".join(row["members"]), color=MUTED, size=10)
            _put(ds, r, 3, c["name"])
            _put(ds, r, 4, _level_formula(r, 0, nlevels, labels, 3, ra, rbounds), center=True)
            _put(ds, r, 5, c["final"], bold=True, center=True, fill=EDIT)
            _put(ds, r, 6, f"=INDEX({rb},MATCH(C{r},{ra},0))", center=True, color=MUTED)
            _put(ds, r, 7, why)
            _put(ds, r, 8, ev, color=MUTED, size=10)
            _fit(ds, r, widths, ["", "", c["name"], "", "", "", why, ev])
            r += 1
        for col in (1, 2, 4, 7, 8):
            _put(ds, r, col, None, fill=STRIPE)
        _put(ds, r, 3, "Team total", bold=True, fill=STRIPE)
        _put(ds, r, 5, f"=SUM(E{start}:E{r - 1})", bold=True, center=True, fill=STRIPE)
        _put(ds, r, 6, f"=SUM(F{start}:F{r - 1})", bold=True, center=True, fill=STRIPE)
        ds.row_dimensions[r].height = 22
        r += 2
    end = max(r, 6)
    for i, w in enumerate(widths, start=1):
        ds.column_dimensions[get_column_letter(i)].width = w
    for name, font_color, fill_color in (
            (labels[0] if labels else "", OK, OK_SOFT),
            (labels[1] if len(labels) > 2 else "", VIOLET, GOOD_SOFT),
            (labels[-1] if labels else "", WARN, WARN_SOFT)):
        if name:
            ds.conditional_formatting.add(f"D5:D{end}", FormulaRule(
                formula=[f'$D5="{name}"'], font=Font(color=font_color, bold=True), fill=_fill(fill_color)))

    # Grades: what the professor asked for. Each score adds up that team's rows on the Details tab.
    _title(gs, title, "Grades for the teams you reviewed. The breakdown behind each score is on the Details tab.")
    _header(gs, [f"Student name", f"Score (out of {rubric.total_points:g})", "Needs your attention", "AI usage"])
    gw = [22, 19, 70, 55]
    g = 5
    for row in rows:
        for name in row["members"]:
            fill = STRIPE if g % 2 == 0 else None
            att, ai = attention_text(row), ai_text(row)
            _put(gs, g, 1, name, bold=True, fill=fill)
            _put(gs, g, 2, f'=SUMIFS(Details!$E$5:$E${end},Details!$A$5:$A${end},"{row["id"]}")',
                 bold=True, size=12, center=True, fill=fill)
            _put(gs, g, 3, att, fill=fill)
            _put(gs, g, 4, ai, fill=fill)
            _fit(gs, g, gw, [name, "", att, ai])
            g += 1
    for i, w in enumerate(gw, start=1):
        gs.column_dimensions[get_column_letter(i)].width = w
    if g > 5:
        gs.conditional_formatting.add(f"D5:D{g - 1}", FormulaRule(
            formula=['LEFT($D5,6)="Review"'], font=Font(color=WARN, bold=True), fill=_fill(WARN_SOFT)))
    wb.properties.creator = "EasyGrade"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_review_excel(rows: list[dict], rubric, path: Path, title: str | None = None) -> dict:
    ok, refused = split_by_guard(rows, rubric)
    build_review_workbook(ok, rubric, path, title)
    return {"written": [r["id"] for r in ok], "refused": refused,
            "students": sum(len(r["members"]) for r in ok)}
