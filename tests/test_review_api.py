"""The review dashboard's Python side: rubric ranges, file matching, notebook checks, the review list and the
guarded exports. No window and no browser; settings are isolated so no API key can be found.
"""
import csv
import json
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from desktop.api import Api
from desktop.review import friendly_flags, original_name, split_members
from desktop.review_export import guard_rubric
from desktop.sample import create_sample_folder
from grader.rubric import Rubric, RubricError, load_rubric

ROOT = Path(__file__).resolve().parents[1]
CAP = ROOT / "rubrics" / "cap3321c_final.json"


class _Isolated(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._env = {k: os.environ.get(k) for k in ("EASYGRADE_CONFIG_DIR", "ANTHROPIC_API_KEY")}
        os.environ["EASYGRADE_CONFIG_DIR"] = str(Path(self.tmp.name) / "config")
        os.environ.pop("ANTHROPIC_API_KEY", None)

    def tearDown(self):
        for k, v in self._env.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        self.tmp.cleanup()

    def api(self, rubric=True):
        api = Api()
        folder = create_sample_folder(Path(self.tmp.name) / "sample")
        self.assertTrue(api.open_folder(str(folder))["ok"])
        if rubric:
            self.assertTrue(api.use_builtin_rubric("cap3321c_final.json")["ok"])
        return api


class RubricRangeTests(unittest.TestCase):
    def test_cap_rubric_loads_with_ranges(self):
        r = load_rubric(CAP)
        self.assertEqual((r.assignment, r.total_points, len(r.criteria)), ("CAP3321C Final Project", 200, 6))
        cleaning = r.criteria[0]
        self.assertEqual([(l.label, l.score, l.floor) for l in cleaning.levels],
                         [("Excellent", 40, 36), ("Good", 32, 28), ("Needs work", 14, 0)])

    def test_a_typed_score_maps_to_the_right_level(self):
        c = load_rubric(CAP).criteria[0]
        for score, label in ((40, "Excellent"), (36, "Excellent"), (35, "Good"), (28, "Good"), (27, "Needs work"),
                             (0, "Needs work")):
            self.assertEqual(c.label_for(score).label, label, score)
        self.assertIsNone(c.label_for(None))

    def test_rubrics_without_ranges_still_work(self):
        c = load_rubric(ROOT / "rubrics" / "project2_regression.json").criteria[0]
        self.assertEqual(c.label_for(11).label, "Proficient")
        self.assertEqual(c.label_for(12).label, "Proficient")
        self.assertEqual(c.label_for(0).label, "Missing")

    def test_ranges_round_trip_and_are_validated(self):
        r = load_rubric(CAP)
        again = Rubric.from_dict(json.loads(r.to_json()))
        self.assertEqual(r.criteria[0].levels, again.criteria[0].levels)
        data = json.loads(r.to_json())
        data["criteria"][0]["levels"][1]["min"] = 99  # above its own score
        with self.assertRaises(RubricError):
            Rubric.from_dict(data)

    def test_guard_ranges_cover_each_level_up_to_the_next(self):
        sec = guard_rubric(load_rubric(CAP))["sections"][0]
        self.assertEqual([(l["level"], l["min_points"]) for l in sec["levels"]],
                         [("Excellent", 36), ("Good", 28), ("Needs work", 0)])
        self.assertEqual(sec["levels"][0]["max_points"], 40)
        self.assertLess(sec["levels"][1]["max_points"], 36)


class HelperTests(unittest.TestCase):
    def test_split_members(self):
        self.assertEqual(split_members("Isabel Moreno & Daniel Park"), ["Isabel Moreno", "Daniel Park"])
        self.assertEqual(split_members("A Lee; B Kim and C Roe"), ["A Lee", "B Kim", "C Roe"])
        self.assertEqual(split_members("Rivera, Ana"), ["Rivera, Ana"])  # a comma is part of the name
        self.assertEqual(split_members("", "1101"), ["1101"])

    def test_original_name_drops_the_lms_prefix(self):
        self.assertEqual(original_name("moreno_1101_88101_Moreno_Park_FinalProject.ipynb"), "Moreno_Park_FinalProject.ipynb")
        self.assertEqual(original_name("final.ipynb"), "final.ipynb")


class RubricApiTests(_Isolated):
    def test_list_and_preview(self):
        api = self.api(rubric=False)
        listing = api.list_rubrics()
        self.assertEqual({r["file"] for r in listing["rubrics"]}, {"cap3321c_final.json", "project2_regression.json"})
        prev = api.preview_builtin_rubric("cap3321c_final.json")
        self.assertTrue(prev["ok"])
        self.assertEqual((prev["total_points"], len(prev["criteria"])), (200, 6))
        self.assertEqual(prev["criteria"][0]["levels"][0], {"label": "Excellent", "score": 40, "min": 36})

    def test_preview_of_a_bad_file_says_why(self):
        bad = Path(self.tmp.name) / "bad.json"
        bad.write_text(json.dumps({"assignment": "x", "criteria": []}), encoding="utf-8")
        res = self.api(rubric=False).preview_rubric(str(bad))
        self.assertFalse(res["ok"])
        self.assertIn("no criteria", res["error"].lower())

    def test_builtin_rubric_names_cannot_escape_the_rubrics_folder(self):
        api = self.api(rubric=False)
        for name in ("../README.md", "..\\README.md", "nope.json"):
            self.assertFalse(api.use_builtin_rubric(name)["ok"], name)

    def test_changing_a_graded_rubric_warns_and_can_be_confirmed(self):
        api = self.api()
        api.grade_simulated("1101")
        data = json.loads(CAP.read_text(encoding="utf-8"))  # same assignment, different wording
        data["criteria"][0]["name"] = "Cleaning"
        changed = Path(self.tmp.name) / "changed.json"
        changed.write_text(json.dumps(data), encoding="utf-8")
        res = api.use_rubric(str(changed))
        self.assertFalse(res["ok"])
        self.assertIn("already has scores", res["error"])
        self.assertTrue(api.use_rubric(str(changed), True)["ok"])


class SampleAndMatchingTests(_Isolated):
    def test_sample_folder_has_the_expected_shape(self):
        folder = create_sample_folder(Path(self.tmp.name) / "s")
        names = {p.name for p in folder.iterdir()}
        self.assertIn("IMG_4821.MOV", names)
        self.assertIn("roster.csv", names)
        self.assertEqual(sum(n.endswith("FinalProject.ipynb") or n.endswith("final.ipynb") for n in names), 9)  # team-008 has none
        with (folder / "roster.csv").open(encoding="utf-8") as f:
            self.assertEqual(len(list(csv.DictReader(f))), 10)

    def test_matches_start_with_three_teams_needing_a_video(self):
        m = self.api().get_matches()
        self.assertEqual((m["needs"], m["matched"]), (3, 6))
        self.assertEqual([u["name"] for u in m["unmatched"]], ["IMG_4821.MOV"])
        status = {t["id"]: t["status"] for t in m["teams"]}
        self.assertEqual(status["1108"], "no_notebook")

    def test_assign_and_no_video_survive_a_rescan_and_a_restart(self):
        api = self.api()
        loose = api.get_matches()["unmatched"][0]["path"]
        m = api.assign_file(loose, "1105")
        self.assertEqual(next(t for t in m["teams"] if t["id"] == "1105")["status"], "by_you")
        m = api.set_no_video("1107")
        m = api.set_no_video("1110")
        self.assertEqual(m["needs"], 0)
        api.scan()
        self.assertEqual(api.get_matches()["needs"], 0)
        api2 = Api()  # a new session on the same folder
        api2.open_folder(str(api._folder))
        self.assertEqual(api2.get_matches()["needs"], 0)
        state = {s["id"]: s for s in api2.get_state()["students"]}
        self.assertTrue(state["1105"]["has_video"])
        self.assertFalse(any("No video found" in f for f in state["1107"]["flags"]))
        # undo
        m = api.assign_file(loose, "")
        self.assertEqual(m["needs"], 1)

    def test_a_file_outside_the_folder_cannot_be_assigned(self):
        api = self.api()
        outside = Path(self.tmp.name) / "x.mp4"
        outside.write_bytes(b"x")
        self.assertFalse(api.assign_file(str(outside), "1105")["ok"])


class NotebookCheckTests(_Isolated):
    def flags(self, api, sid):
        res = api.analyze_notebook(sid)
        self.assertTrue(res["ok"], res)
        return res["flags"]

    def test_clean_team_has_no_flags(self):
        self.assertEqual(self.flags(self.api(), "1101"), [])

    def test_wrong_file_name_and_run_order(self):
        flags = self.flags(self.api(), "1104")
        self.assertTrue(any("tomas-grace final.ipynb" in f and "Lastname1_Lastname2" in f for f in flags), flags)
        self.assertTrue(any(f.startswith("Cells were run out of order") for f in flags), flags)

    def test_error_cell_is_reported_with_a_one_based_cell_number(self):
        flags = self.flags(self.api(), "1106")
        self.assertIn("The notebook stops with an error (cell 5)", flags)

    def test_empty_notebook_and_missing_names(self):
        flags = self.flags(self.api(), "1109")
        self.assertIn("The notebook is empty", flags)
        self.assertTrue(any("header is missing" in f.lower() for f in flags), flags)

    def test_no_notebook_is_skipped_and_flags_are_saved_once(self):
        api = self.api()
        self.assertTrue(api.analyze_notebook("1108")["skipped"])
        api.analyze_notebook("1104")
        api.analyze_notebook("1104")  # twice: no duplicates
        flags = {s["id"]: s["flags"] for s in api.get_state()["students"]}["1104"]
        self.assertEqual(len(flags), len(set(flags)))

    def test_submission_rules_only_apply_to_the_cap_rubric(self):
        api = self.api(rubric=False)
        self.assertTrue(api.use_builtin_rubric("project2_regression.json")["ok"])
        flags = self.flags(api, "1104")
        self.assertFalse(any("Lastname1_Lastname2" in f for f in flags), flags)


class ReviewListTests(_Isolated):
    def test_before_grading_every_criterion_needs_a_score(self):
        rows = {r["id"]: r for r in self.api().get_review()["rows"]}
        row = rows["1101"]
        self.assertFalse(row["complete"])
        self.assertIsNone(row["total"])
        self.assertEqual(row["overall"], "Not graded yet.")
        self.assertEqual(len([c for c in row["check"] if "needs your score" in c]), 6)

    def test_simulated_grading_never_invents_a_video_score(self):
        api = self.api()
        for sid in ("1101", "1105"):  # 1105 has no video
            self.assertTrue(api.grade_simulated(sid)["ok"])
        rows = {r["id"]: r for r in api.get_review()["rows"]}
        self.assertTrue(rows["1101"]["complete"])
        gaps = [c["id"] for c in rows["1105"]["criteria"] if c["final"] is None]
        self.assertEqual(gaps, ["presentation"])
        self.assertTrue(any("needs your score (no video)" in c for c in rows["1105"]["check"]))
        self.assertIn("so far", rows["1105"]["overall"])
        self.assertIsNotNone(rows["1105"]["total"])  # points so far

    def test_typed_scores_get_a_level_and_change_tracking(self):
        api = self.api()
        api.grade_simulated("1101")
        self.assertTrue(api.set_score("1101", "data_cleaning", 38)["ok"])
        row = {r["id"]: r for r in api.get_review()["rows"]}["1101"]
        c = next(c for c in row["criteria"] if c["id"] == "data_cleaning")
        self.assertEqual((c["final"], c["level"], c["changed"]), (38, "Excellent", True))
        self.assertNotEqual(row["changed_by"], 0)
        self.assertFalse(api.set_score("1101", "data_cleaning", 41)["ok"])  # above the 40 points
        self.assertFalse(api.set_score("1101", "data_cleaning", -1)["ok"])

    def test_strongest_and_weakest_follow_the_scores(self):
        api = self.api()
        api.grade_simulated("1101")
        for cid, v in (("data_cleaning", 40), ("collaboration", 0)):
            api.set_score("1101", cid, v)
        row = {r["id"]: r for r in api.get_review()["rows"]}["1101"]
        self.assertEqual(row["strongest"]["name"], "Data Cleaning & Preparation")
        self.assertEqual(row["weakest"]["name"], "Collaboration & Formatting")

    def test_approving_needs_every_score(self):
        api = self.api()
        api.grade_simulated("1105")  # presentation missing
        self.assertFalse(api.approve("1105")["ok"])
        api.set_score("1105", "presentation", 30)
        self.assertTrue(api.approve("1105")["ok"])


class ExportTests(_Isolated):
    def graded(self, approve=("1101", "1102")):
        api = self.api()
        for sid in ("1101", "1102", "1103"):
            api.grade_simulated(sid)
        for sid in approve:
            self.assertTrue(api.approve(sid)["ok"])
        return api

    def test_only_reviewed_teams_reach_the_files(self):
        api = self.graded()
        x = api.export_review("excel")
        self.assertEqual((x["teams"], x["rows"], x["refused"]), (2, 4, {}))
        self.assertGreater(x["left_out"], 0)
        c = api.export_review("csv")
        with open(c["file"], encoding="utf-8", newline="") as f:
            rows = list(csv.reader(f))
        self.assertEqual(rows[0], ["Student name", "Score (out of 200)", "Needs your attention", "AI usage"])
        self.assertEqual({r[0] for r in rows[1:]}, {"Isabel Moreno", "Daniel Park", "Camila Reyes", "Omar Haddad"})
        self.assertEqual(rows[1][1], rows[2][1])  # both members of a team get the team's grade
        self.assertNotIn(".", rows[1][1])  # 108, not 108.0
        self.assertTrue((Path(c["file"]).parent / "export_log.jsonl").is_file())

    def test_nothing_is_written_when_nothing_is_reviewed(self):
        api = self.graded(approve=())
        x = api.export_review("csv")
        self.assertEqual((x["teams"], x["rows"]), (0, 0))
        with open(x["file"], encoding="utf-8") as f:
            self.assertEqual(len(f.read().splitlines()), 1)  # header only

    def test_excel_has_the_three_tabs_and_working_formulas(self):
        api = self.graded()
        path = api.export_review("excel")["file"]
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Grades", "Details", "Rubric"])
        grades = wb["Grades"]
        self.assertEqual([c.value for c in grades[4][:4]], ["Student name", "Score (out of 200)", "Needs your attention", "AI usage"])
        self.assertIn('SUMIFS(Details!', grades["B5"].value)
        self.assertIn('"1102"', grades["B5"].value)  # exact team id, not a name match
        details = wb["Details"]
        levels = [r[3].value for r in details.iter_rows(min_row=5) if r[3].value]
        self.assertTrue(all(v.startswith("=IF(E") for v in levels))
        rubric = wb["Rubric"]
        self.assertEqual([c.value for c in rubric[4][:5]], ["Category", "Points", "Excellent from", "Good from", "Needs work"])
        self.assertEqual(rubric["C5"].value, 36)
        self.assertEqual(rubric["B11"].value, "=SUM(B5:B10)")

    def test_the_guard_refuses_a_reviewed_team_with_a_missing_score(self):
        from desktop import review_export as rx

        api = self.graded()
        rows = [r for r in api.review_rows() if r["status"] == "approved"]
        rows[0]["criteria"][0]["final"] = None  # as if the store let it through
        ok, refused = rx.split_by_guard(rows, api._rubric)
        self.assertEqual(len(ok), 1)
        self.assertIn(rows[0]["id"], refused)

    def test_spreadsheet_formulas_in_names_are_neutralised(self):
        from desktop import review_export as rx

        api = self.graded()
        rows = [r for r in api.review_rows() if r["status"] == "approved"]
        rows[0]["members"] = ["=HYPERLINK(\"http://x\")", "Normal Name"]
        res = rx.export_review_csv(rows, api._rubric, Path(self.tmp.name) / "out.csv")
        self.assertEqual(res["rows_written"], 4)
        with open(Path(self.tmp.name) / "out.csv", encoding="utf-8", newline="") as f:
            for row in csv.reader(f):
                self.assertFalse(row[0].startswith("="), row[0])

    def test_legacy_export_still_works(self):
        api = self.graded()
        res = api.export_all()
        self.assertTrue(res["ok"], res)
        self.assertTrue((Path(res["folder"]) / "grades.xlsx").is_file())


class FriendlyFlagTests(unittest.TestCase):
    def test_messages_are_plain_sentences(self):
        from easygrade.pipeline.notebook.config import load_config
        from easygrade.pipeline.notebook.models import Checks, ExecutionOrder, NotebookCells, make_flag

        res = NotebookCells(
            schema_version="1", team_id="t", produced_by="x", header={"found": True, "fields": {}, "rules_ok": False},
            filename_ok=False,
            checks=Checks(empty=False, execution_order=ExecutionOrder(strictly_increasing=False, out_of_order_cells=[2, 4]),
                          error_cells=[3, 6]),
            flags=[make_flag("EXEC_ORDER"), make_flag("ERROR_OUTPUT"),
                   make_flag("HEADER_RULES", "Header is missing required fields: date, section.")])
        out = friendly_flags(res, "f.ipynb", load_config())
        self.assertEqual(out, ["Cells were run out of order (cells 3, 5)", "Cells with errors: 4, 7",
                               "The header is missing: date, section"])


if __name__ == "__main__":
    unittest.main()
