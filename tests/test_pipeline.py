"""Run with:  python -m unittest discover tests"""
import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from grader import GradeStore, RubricError, load_rubric, parse_grader_output, write_rubric_template
from grader.cli import SAMPLE_RUBRIC, main
from grader.scoring import GraderOutputError, SimulatedGrader
from grader.store import StoreError


class RubricTests(unittest.TestCase):
    def test_sample_rubric_loads(self):
        r = load_rubric(SAMPLE_RUBRIC)
        self.assertEqual(r.total_points, 100)
        self.assertEqual(len(r.criteria), 7)

    def test_excel_round_trip(self):
        r = load_rubric(SAMPLE_RUBRIC)
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "r.xlsx"
            write_rubric_template(p, r)
            r2 = load_rubric(p)
            self.assertEqual(r.to_dict()["criteria"], r2.to_dict()["criteria"])
            self.assertEqual(r2.assignment, r.assignment)

    def test_invalid_rubric_lists_all_problems(self):
        data = {"assignment": "X", "criteria": [
            {"id": "a", "name": "A", "source": "audio", "points": 10, "levels": [{"score": 8, "label": "hi"}]},
            {"id": "a", "name": "A2", "source": "video", "points": 5, "levels": [{"score": 5, "label": "hi"}]},
        ]}
        from grader import Rubric
        with self.assertRaises(RubricError) as cm:
            Rubric.from_dict(data)
        text = str(cm.exception)
        self.assertIn("source must be one of", text)
        self.assertIn("top level scores 8", text)
        self.assertIn("duplicate id", text)


class ScoringTests(unittest.TestCase):
    def setUp(self):
        self.r = load_rubric(SAMPLE_RUBRIC)

    def test_rejects_score_not_in_levels(self):
        out = SimulatedGrader().grade(self.r, "x")
        out["criteria"]["modeling"]["score"] = 17
        with self.assertRaises(GraderOutputError):
            parse_grader_output(self.r, out)

    def test_rejects_missing_criterion(self):
        out = SimulatedGrader().grade(self.r, "x")
        del out["criteria"]["eda"]
        with self.assertRaises(GraderOutputError):
            parse_grader_output(self.r, out)


class StoreTests(unittest.TestCase):
    def test_override_survives_regrade_and_approve_requires_all(self):
        r = load_rubric(SAMPLE_RUBRIC)
        with tempfile.TemporaryDirectory() as d, GradeStore(Path(d) / "g.db") as st:
            aid = st.save_assignment(r)
            sid = st.upsert_submission(aid, "1", "Test Student")
            with self.assertRaises(StoreError):
                st.approve(sid, r)
            res, fb, _ = parse_grader_output(r, SimulatedGrader().grade(r, "1"))
            st.record_ai_results(sid, r, res, fb)
            st.override_score(sid, r, "modeling", 1.5, "manual")
            st.record_ai_results(sid, r, res, fb)  # regrade must not wipe the override
            self.assertEqual(st.scores(sid)["modeling"]["final_score"], 1.5)
            with self.assertRaises(StoreError):
                st.override_score(sid, r, "modeling", 99)
            st.approve(sid, r)
            self.assertEqual(st.submission(aid, "1").status, "approved")


class DemoEndToEnd(unittest.TestCase):
    def test_demo_exports(self):
        with tempfile.TemporaryDirectory() as d:
            self.assertEqual(main(["demo", d]), 0)
            out = Path(d)
            wb = load_workbook(out / "grades.xlsx")
            self.assertEqual(wb.sheetnames, ["Summary", "Detail", "Feedback", "Flags", "Rubric"])
            ws = wb["Summary"]
            ids = [ws.cell(r, 1).value for r in range(2, 12)]
            self.assertEqual(len([i for i in ids if i]), 10)
            with (out / "gradebook_simple.csv").open() as f:
                rows = list(csv.DictReader(f))
            self.assertEqual(len(rows), 7)  # only approved
            self.assertTrue(all(0 <= float(r["score"]) <= 100 for r in rows))
            with (out / "gradebook_for_lms.csv").open() as f:
                lms = list(csv.reader(f))
            col = lms[0].index("Project 2 - Housing Price Regression (4411)")
            filled = [row for row in lms[2:] if row[col]]
            self.assertEqual(len(filled), 7)
            self.assertEqual(lms[1][col], "100")  # points-possible row untouched
            self.assertEqual(len(list((out / "feedback").glob("*.txt"))), 9)


if __name__ == "__main__":
    unittest.main()
